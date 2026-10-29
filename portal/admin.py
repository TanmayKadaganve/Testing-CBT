import csv
import datetime
import io

from django import forms
from django.contrib import admin, messages
from django.http import HttpResponse
from django.shortcuts import render, redirect
from django.urls import path
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from django.db import transaction
from portal.services.language_registry import LanguageRegistry
from portal.services.spreadsheet import escape_formula_cells

from .models import (
    Student, Exam, Section, Question, Option, Attempt, Answer, Feedback, CodingQuestion, TestCase, CodeDraft, CodeSubmission, ExecutionJob, CodingQuestionStarterCode
)

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def fmt(dt):
    return timezone.localtime(dt).strftime("%d-%m-%Y %H:%M:%S") if dt else ""


def xlsx_response(wb, filename):
    escape_formula_cells(wb)
    response = HttpResponse(content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    escape_formula_cells(wb).save(response)
    return response


# ---------- Results export ----------

@admin.action(description="Download results (Excel)")
def download_results(modeladmin, request, queryset):
    if queryset.count() != 1:
        modeladmin.message_user(request, "Please select exactly one exam.", level=messages.ERROR)
        return
    exam = queryset.first()

    sections = list(exam.sections.order_by("order"))
    questions = list(
        Question.objects.filter(section__exam=exam)
        .select_related("section")
        .order_by("section__order", "order")
    )
    students = list(exam.students.order_by("usn"))
    attempts = {a.student_id: a for a in Attempt.objects.filter(exam=exam)}
    answers = {}
    for ans in Answer.objects.filter(attempt__exam=exam).select_related("selected_option"):
        answers[(ans.attempt_id, ans.question_id)] = ans

    max_marks = sum(q.marks for q in questions)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(
        ["USN", "Name", "Status", "Started", "Submitted"]
        + [f"{s.title} score" for s in sections]
        + ["Total score", "Max marks"]
    )

    detail = wb.create_sheet("Answers")
    detail.append(["USN", "Name"] + [f"{q.section.title} Q{q.order}" for q in questions])

    for sheet in (summary, detail):
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    for st in students:
        att = attempts.get(st.id)
        if not att:
            status = "Not attempted"
        elif att.submitted_at:
            status = "Submitted"
        else:
            status = "In progress"

        section_scores = {s.id: 0 for s in sections}
        chosen = []
        for q in questions:
            ans = answers.get((att.id, q.id)) if att else None
            opt = ans.selected_option if ans else None
            chosen.append(opt)
            if opt and opt.is_correct:
                section_scores[q.section_id] += q.marks

        summary.append(
            [
                st.usn,
                st.name,
                status,
                fmt(att.started_at) if att else "",
                fmt(att.submitted_at) if att else "",
            ]
            + [section_scores[s.id] for s in sections]
            + [sum(section_scores.values()), max_marks]
        )

        detail.append([st.usn, st.name] + [opt.text if opt else "-" for opt in chosen])
        row = detail.max_row
        for i, opt in enumerate(chosen):
            if opt:
                detail.cell(row=row, column=3 + i).fill = green if opt.is_correct else red

    for sheet in (summary, detail):
        for col in sheet.columns:
            longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(longest + 2, 10), 40)

    return xlsx_response(wb, f"results_exam_{exam.id}.xlsx")


# ---------- Student import ----------

class ImportForm(forms.Form):
    file = forms.FileField(label="CSV or Excel file")
    exam = forms.ModelChoiceField(
        queryset=Exam.objects.all(),
        required=False,
        label="Also register these students for this exam",
    )


def read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(upload.read()), data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    raise ValueError("Please upload a .csv or .xlsx file.")


def find_col(headers, keys, avoid=()):
    for i, h in enumerate(headers):
        if any(k in h for k in keys) and not any(a in h for a in avoid):
            return i
    return None


def parse_dob(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    s = str(value or "").strip()
    if not s:
        return None
    formats = [
        "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d",
        "%d-%m-%y", "%d/%m/%y", "%d %b %Y", "%d-%b-%Y", "%d %B %Y",
    ]
    for f in formats:
        try:
            return datetime.datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def import_students(rows, exam):
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("The file is empty.")

    headers = [str(h or "").strip().lower() for h in rows[0]]
    usn_i = find_col(headers, ["usn", "roll no", "reg no"])
    name_i = find_col(headers, ["name"], avoid=["usn"])
    dob_i = find_col(headers, ["dob", "birth"])
    email_i = find_col(headers, ["email", "mail"])

    missing = [label for label, idx in (("USN", usn_i), ("Name", name_i), ("DOB", dob_i)) if idx is None]
    if missing:
        raise ValueError(
            "Could not find column(s): " + ", ".join(missing)
            + ". Headings found: " + ", ".join(h for h in headers if h)
        )

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    created = updated = 0
    problems = []
    students = []

    for line_no, row in enumerate(rows[1:], start=2):
        usn = str(cell(row, usn_i) or "").strip().upper()
        name = str(cell(row, name_i) or "").strip()
        raw_dob = cell(row, dob_i)
        dob = parse_dob(raw_dob)

        if not usn:
            problems.append(f"Row {line_no}: USN is missing")
            continue
        if not name:
            problems.append(f"Row {line_no} ({usn}): name is missing")
            continue
        if not dob:
            problems.append(f"Row {line_no} ({usn}): could not read DOB '{raw_dob}'")
            continue

        defaults = {"name": name, "dob": dob}
        if email_i is not None:
            defaults["email"] = str(cell(row, email_i) or "").strip()

        student, was_created = Student.objects.update_or_create(usn=usn, defaults=defaults)
        if was_created:
            created += 1
        else:
            updated += 1
        students.append(student)

    if exam and students:
        exam.students.add(*students)
    return created, updated, problems

class QuestionImportForm(forms.Form):
    file = forms.FileField(label="Excel or CSV file")
    exam = forms.ModelChoiceField(queryset=Exam.objects.all(), label="Import into exam")


class CodingQuestionImportForm(forms.Form):
    file = forms.FileField(label="Coding questions workbook (.xlsx)")
    exam = forms.ModelChoiceField(queryset=Exam.objects.all(), label="Import into exam")


QNO_HEADERS = {"q no", "qno", "q.no", "q. no", "question no", "question number", "no", "sl no", "s.no"}
TEXT_HEADERS = {"question", "questions", "question text"}


def find_header(headers, names):
    for i, h in enumerate(headers):
        if h in names:
            return i
    return None


def option_columns(headers):
    found = []
    for letter in "abcdef":
        names = {f"option {letter}", f"option{letter}", f"opt {letter}", f"option_{letter}", letter}
        idx = find_header(headers, names)
        if idx is not None:
            found.append((letter.upper(), idx))
    return found


def to_int(value):
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def import_questions(rows, exam):
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("The file is empty.")

    headers = [str(h or "").strip().lower() for h in rows[0]]
    section_i = find_col(headers, ["section"])
    qno_i = find_header(headers, QNO_HEADERS)
    text_i = find_header(headers, TEXT_HEADERS)
    if text_i is None:
        text_i = find_col(headers, ["question"], avoid=["no", "number"])
    correct_i = find_col(headers, ["correct", "answer"])
    marks_i = find_col(headers, ["mark"])
    options = option_columns(headers)

    missing = [
        label
        for label, idx in (
            ("Section", section_i),
            ("Q No", qno_i),
            ("Question", text_i),
            ("Correct", correct_i),
        )
        if idx is None
    ]
    if len(options) < 2:
        missing.append("Option A and Option B (at least two option columns)")
    if missing:
        raise ValueError(
            "Could not find: " + ", ".join(missing)
            + ". Headings found: " + ", ".join(h for h in headers if h)
        )

    sections = {s.title.strip().lower(): s for s in exam.sections.all()}

    def cell(row, i):
        return row[i] if i is not None and i < len(row) else None

    created = updated = 0
    problems = []

    for line_no, row in enumerate(rows[1:], start=2):
        title = str(cell(row, section_i) or "").strip()
        section = sections.get(title.lower())
        if not section:
            available = ", ".join(s.title for s in sections.values()) or "none"
            problems.append(
                f"Row {line_no}: section '{title}' not found in this exam (available: {available})"
            )
            continue

        qno = to_int(cell(row, qno_i))
        text = str(cell(row, text_i) or "").strip()
        if qno is None or not text:
            problems.append(f"Row {line_no}: Q No or Question is missing")
            continue

        opts = []
        for letter, idx in options:
            val = cell(row, idx)
            if val not in (None, ""):
                opts.append((letter, str(val).strip()))
        if len(opts) < 2:
            problems.append(f"Row {line_no} (Q{qno}): needs at least two options")
            continue

        correct_value = str(cell(row, correct_i) or "").strip()
        correct = correct_value.upper()
        option_letters = [letter for letter, _ in opts]
        if correct not in option_letters:
            correct = next(
                (letter for letter, option_text in opts
                 if option_text.strip().casefold() == correct_value.casefold()),
                "",
            )
        if not correct:
            problems.append(
                f"Row {line_no} (Q{qno}): Correct value '{correct_value.upper()}' does not match an option"
            )
            continue
        
        marks = 1
        if marks_i is not None and cell(row, marks_i) not in (None, ""):
            marks = to_int(cell(row, marks_i))
            if not marks or marks < 1:
                problems.append(f"Row {line_no} (Q{qno}): marks must be a whole number")
                continue

        existing = Question.objects.filter(section=section, order=qno).first()
        if existing and Answer.objects.filter(question=existing).exists():
            problems.append(
                f"Row {line_no} (Q{qno}): skipped, students have already answered this question"
            )
            continue

        with transaction.atomic():
            if existing:
                existing.text = text
                existing.marks = marks
                existing.save()
                existing.options.all().delete()
                question = existing
                updated += 1
            else:
                question = Question.objects.create(section=section, order=qno, text=text, marks=marks)
                created += 1
            for letter, opt_text in opts:
                Option.objects.create(
                    question=question, text=opt_text, is_correct=(letter == correct)
                )

    return created, updated, problems


def question_template_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(["Section", "Q No", "Question", "Option A", "Option B", "Option C", "Option D", "Correct", "Marks"])
    ws.append(["MCQ", 1, "Which keyword is used to create an object in Java?", "new", "create", "make", "alloc", "A", 5])
    ws.append(["MCQ", 2, "What is the output of 10 / 3 in Java?", "3.33", "3", "4", "3.0", "B", 5])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=300, min_col=4, max_col=7):
        for c in row:
            c.number_format = "@"
    widths = {"A": 14, "B": 7, "C": 60, "D": 22, "E": 22, "F": 22, "G": 22, "H": 9, "I": 7}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    notes = wb.create_sheet("Read me")
    for line in [
        "Section: must match a section title of the exam exactly (for example MCQ).",
        "Q No: the question's number within its section. Importing the same Q No again updates that question.",
        "Question: plain text. Do not type the number again at the start.",
        "Options: leave unused options blank. At least two are needed. Option cells are formatted as text.",
        "Correct: the letter of the right option (A, B, C, D...).",
        "Marks: optional, defaults to 1.",
        "Questions that students have already answered are never changed by an import.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    return wb


def coding_question_template_workbook():
    wb = Workbook()
    questions = wb.active
    questions.title = "Coding Questions"
    questions.append([
        "Section", "Q No", "Title", "Problem statement", "Input format", "Output format",
        "Constraints", "Sample input", "Sample output", "Time limit seconds", "Marks",
    ])
    questions.append([
        "Coding", 1, "Add Two Numbers",
        "Read two integers, add them, and print their sum.",
        "Two integers separated by a space.", "Print their sum.",
        "-10^9 <= A, B <= 10^9", "10 20", "30", 2, 10,
    ])

    test_cases = wb.create_sheet("Test Cases")
    test_cases.append(["Section", "Q No", "Input data", "Expected output", "Is sample", "Marks"])
    test_cases.append(["Coding", 1, "10 20", "30", "Yes", 1])
    test_cases.append(["Coding", 1, "5 7", "12", "No", 1])
    test_cases.append(["Coding", 1, "-10 4", "-6", "No", 1])

    notes = wb.create_sheet("Read me")
    for line in [
        "Use the Coding Questions sheet for one row per coding question.",
        "Use the Test Cases sheet for one row per test case; repeat Section and Q No for each case.",
        "Section must match a CODING section in the selected exam exactly.",
        "Q No identifies the question within its section. Re-importing updates the question and its test cases.",
        "Is sample accepts Yes/No, True/False, or 1/0. Hidden test cases should be No.",
        "Time limit seconds and Marks are optional; they default to 2 seconds and 10 marks.",
        "Questions with existing code drafts or submissions are skipped to preserve student work and grading history.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110

    widths = {
        "A": 18, "B": 8, "C": 26, "D": 58, "E": 38, "F": 38,
        "G": 38, "H": 22, "I": 22, "J": 20, "K": 10,
    }
    for column, width in widths.items():
        questions.column_dimensions[column].width = width
    test_widths = {"A": 18, "B": 8, "C": 38, "D": 38, "E": 14, "F": 10}
    for column, width in test_widths.items():
        test_cases.column_dimensions[column].width = width
    for sheet in (questions, test_cases):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    return wb


def read_coding_workbook(upload):
    if not upload.name.lower().endswith(".xlsx"):
        raise ValueError("Please upload the coding questions .xlsx template.")
    wb = load_workbook(io.BytesIO(upload.read()), data_only=True)
    sheets = {ws.title.strip().casefold(): ws for ws in wb.worksheets}
    if "coding questions" not in sheets or "test cases" not in sheets:
        raise ValueError("The workbook must contain 'Coding Questions' and 'Test Cases' sheets.")
    return (
        [list(row) for row in sheets["coding questions"].iter_rows(values_only=True)],
        [list(row) for row in sheets["test cases"].iter_rows(values_only=True)],
    )


def import_coding_questions(question_rows, test_case_rows, exam):
    question_rows = [r for r in question_rows if r and any(c not in (None, "") for c in r)]
    test_case_rows = [r for r in test_case_rows if r and any(c not in (None, "") for c in r)]
    if not question_rows or not test_case_rows:
        raise ValueError("Both workbook sheets need a header row and data rows.")

    question_headers = [str(h or "").strip().lower() for h in question_rows[0]]
    test_headers = [str(h or "").strip().lower() for h in test_case_rows[0]]

    question_cols = {
        "section": find_col(question_headers, ["section"]),
        "qno": find_header(question_headers, QNO_HEADERS),
        "title": find_header(question_headers, {"title", "question title"}),
        "statement": find_col(question_headers, ["problem statement", "statement", "problem"]),
        "input_format": find_col(question_headers, ["input format"]),
        "output_format": find_col(question_headers, ["output format"]),
        "constraints": find_col(question_headers, ["constraint"]),
        "sample_input": find_col(question_headers, ["sample input"]),
        "sample_output": find_col(question_headers, ["sample output"]),
        "time_limit": find_col(question_headers, ["time limit"]),
        "marks": find_col(question_headers, ["mark"]),
    }
    test_cols = {
        "section": find_col(test_headers, ["section"]),
        "qno": find_header(test_headers, QNO_HEADERS),
        "input_data": find_col(test_headers, ["input data", "input"]),
        "expected_output": find_col(test_headers, ["expected output", "output"]),
        "is_sample": find_col(test_headers, ["is sample", "sample"]),
        "marks": find_col(test_headers, ["mark"]),
    }

    required_question_cols = [("Section", "section"), ("Q No", "qno"), ("Title", "title"), ("Problem statement", "statement")]
    missing = [label for label, key in required_question_cols if question_cols[key] is None]
    required_test_cols = [("Section", "section"), ("Q No", "qno"), ("Input data", "input_data"), ("Expected output", "expected_output")]
    missing.extend(f"Test Cases: {label}" for label, key in required_test_cols if test_cols[key] is None)
    if missing:
        raise ValueError("Could not find required column(s): " + ", ".join(missing))

    sections = {s.title.strip().casefold(): s for s in exam.sections.all()}

    def cell(row, cols, key):
        index = cols.get(key)
        return row[index] if index is not None and index < len(row) else None

    def text_value(value):
        if value is None:
            return ""
        return str(value).strip()

    def parse_sample(value):
        if isinstance(value, bool):
            return value
        normalized = str(value or "").strip().casefold()
        if normalized in ("", "no", "false", "0", "hidden"):
            return False
        if normalized in ("yes", "true", "1", "sample"):
            return True
        return None

    test_cases_by_question = {}
    invalid_case_keys = set()
    problems = []
    for line_no, row in enumerate(test_case_rows[1:], start=2):
        section_title = text_value(cell(row, test_cols, "section"))
        qno = to_int(cell(row, test_cols, "qno"))
        key = (section_title.casefold(), qno) if qno is not None else None
        if not section_title or qno is None or qno < 1:
            problems.append(f"Test Cases row {line_no}: Section or a positive Q No is missing/invalid")
            if key:
                invalid_case_keys.add(key)
            continue

        raw_sample = cell(row, test_cols, "is_sample")
        is_sample = parse_sample(raw_sample)
        marks_value = cell(row, test_cols, "marks")
        marks = to_int(marks_value) if marks_value not in (None, "") else 1
        if is_sample is None:
            problems.append(f"Test Cases row {line_no} (Q{qno}): Is sample must be Yes or No")
            invalid_case_keys.add(key)
            continue
        if marks is None or marks < 1:
            problems.append(f"Test Cases row {line_no} (Q{qno}): marks must be a positive whole number")
            invalid_case_keys.add(key)
            continue

        test_cases_by_question.setdefault(key, []).append({
            "input_data": text_value(cell(row, test_cols, "input_data")),
            "expected_output": text_value(cell(row, test_cols, "expected_output")),
            "is_sample": is_sample,
            "marks": marks,
        })

    seen = set()
    question_keys = set()
    created = updated = 0
    for line_no, row in enumerate(question_rows[1:], start=2):
        section_title = text_value(cell(row, question_cols, "section"))
        qno = to_int(cell(row, question_cols, "qno"))
        key = (section_title.casefold(), qno) if section_title and qno is not None else None
        if key:
            question_keys.add(key)
        if not section_title or qno is None or qno < 1:
            problems.append(f"Coding Questions row {line_no}: Section or a positive Q No is missing/invalid")
            continue
        if key in seen:
            problems.append(f"Coding Questions row {line_no} (Q{qno}): duplicate Section and Q No")
            continue
        seen.add(key)

        section = sections.get(section_title.casefold())
        if not section:
            problems.append(f"Coding Questions row {line_no}: section '{section_title}' not found in this exam")
            continue
        if section.kind != "CODING":
            problems.append(f"Coding Questions row {line_no} (Q{qno}): section '{section_title}' is not a coding section")
            continue

        title = text_value(cell(row, question_cols, "title"))
        statement = text_value(cell(row, question_cols, "statement"))
        if not title or not statement:
            problems.append(f"Coding Questions row {line_no} (Q{qno}): Title and Problem statement are required")
            continue

        time_value = cell(row, question_cols, "time_limit")
        time_limit = to_int(time_value) if time_value not in (None, "") else 2
        marks_value = cell(row, question_cols, "marks")
        marks = to_int(marks_value) if marks_value not in (None, "") else 10
        if time_limit is None or time_limit < 1:
            problems.append(f"Coding Questions row {line_no} (Q{qno}): time limit must be a positive whole number")
            continue
        if marks is None or marks < 1:
            problems.append(f"Coding Questions row {line_no} (Q{qno}): marks must be a positive whole number")
            continue

        cases = test_cases_by_question.get(key, [])
        if key in invalid_case_keys:
            continue
        sample_input = text_value(cell(row, question_cols, "sample_input"))
        sample_output = text_value(cell(row, question_cols, "sample_output"))
        if (sample_input or sample_output) and not any(tc["is_sample"] for tc in cases):
            cases.append({
                "input_data": sample_input,
                "expected_output": sample_output,
                "is_sample": True,
                "marks": 1,
            })
        if not cases:
            problems.append(f"Coding Questions row {line_no} (Q{qno}): add at least one test case")
            continue
        if not sample_input and not sample_output:
            sample_case = next((tc for tc in cases if tc["is_sample"]), None)
            if sample_case:
                sample_input = sample_case["input_data"]
                sample_output = sample_case["expected_output"]

        existing = CodingQuestion.objects.filter(section=section, order=qno).first()
        if existing and (
            CodeSubmission.objects.filter(question=existing).exists()
            or CodeDraft.objects.filter(question=existing).exists()
        ):
            problems.append(f"Coding Questions row {line_no} (Q{qno}): skipped because students have drafts or submissions")
            continue

        defaults = {
            "title": title,
            "statement": statement,
            "input_format": text_value(cell(row, question_cols, "input_format")),
            "output_format": text_value(cell(row, question_cols, "output_format")),
            "constraints": text_value(cell(row, question_cols, "constraints")),
            "sample_input": sample_input,
            "sample_output": sample_output,
            "time_limit_seconds": time_limit,
            "marks": marks,
        }
        with transaction.atomic():
            if existing:
                for field, value in defaults.items():
                    setattr(existing, field, value)
                existing.save()
                existing.testcases.all().delete()
                question = existing
                updated += 1
            else:
                question = CodingQuestion.objects.create(section=section, order=qno, **defaults)
                created += 1
            TestCase.objects.bulk_create([
                TestCase(question=question, **test_case) for test_case in cases
            ])

    for key in test_cases_by_question:
        if key not in question_keys:
            problems.append(f"Test Cases: no matching coding question found for section '{key[0]}' Q{key[1]}")

    return created, updated, problems


def style_student_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "DD-MM-YYYY"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 30


# ---------- Admin classes ----------

@admin.action(description="Reset login (allow login from another PC)")
def reset_login(modeladmin, request, queryset):
    count = queryset.update(active_token="")
    modeladmin.message_user(request, f"Login reset for {count} student(s).")


class SectionInlineForm(forms.ModelForm):
    class Meta:
        model = Section
        fields = "__all__"
        widgets = {
            "start_time": forms.SplitDateTimeWidget(
                date_format="%Y-%m-%d",
                time_format="%H:%M",
                date_attrs={"type": "date"},
                time_attrs={"type": "time"},
            ),
            "end_time": forms.SplitDateTimeWidget(
                date_format="%Y-%m-%d",
                time_format="%H:%M",
                date_attrs={"type": "date"},
                time_attrs={"type": "time"},
            ),
        }


class SectionInline(admin.TabularInline):
    model = Section
    form = SectionInlineForm
    extra = 1


class OptionInline(admin.TabularInline):
    model = Option
    extra = 4


class ExamSectionSelect(forms.Select):
    def __init__(self, sections, *args, **kwargs):
        self.section_exam_ids = {str(section.pk): section.exam_id for section in sections}
        super().__init__(*args, **kwargs)

    def create_option(self, name, value, label, selected, index, subindex=None, attrs=None):
        option = super().create_option(name, value, label, selected, index, subindex, attrs)
        option_value = getattr(value, "value", value)
        exam_id = self.section_exam_ids.get(str(option_value))
        if exam_id:
            option["attrs"]["data-exam-id"] = exam_id
        return option


class ExamSectionAdminFormMixin:
    section_kind = None

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["exam"].queryset = Exam.objects.filter(
            sections__kind=self.section_kind
        ).distinct()
        sections = Section.objects.filter(kind=self.section_kind).select_related("exam")
        self.fields["section"].queryset = sections
        self.fields["section"].widget = ExamSectionSelect(sections)
        self.fields["section"].widget.choices = self.fields["section"].choices
        if self.instance and self.instance.pk:
            self.initial.setdefault("exam", self.instance.section.exam_id)

    def clean(self):
        cleaned_data = super().clean()
        exam = cleaned_data.get("exam")
        section = cleaned_data.get("section")
        if exam and section:
            if section.exam_id != exam.id:
                self.add_error("section", "Select a section from the selected exam.")
            elif section.kind != self.section_kind:
                self.add_error("section", "Select a section of the correct question type.")
        return cleaned_data


class QuestionAdminForm(ExamSectionAdminFormMixin, forms.ModelForm):
    section_kind = "MCQ"
    exam = forms.ModelChoiceField(queryset=Exam.objects.none(), label="Exam")

    class Meta:
        model = Question
        fields = "__all__"


class CodingQuestionAdminForm(ExamSectionAdminFormMixin, forms.ModelForm):
    section_kind = "CODING"
    exam = forms.ModelChoiceField(queryset=Exam.objects.none(), label="Exam")

    class Meta:
        model = CodingQuestion
        fields = "__all__"


@admin.register(Student)
class StudentAdmin(admin.ModelAdmin):
    list_display = ("usn", "name", "dob", "is_active", "logged_in", "last_login_ip", "last_login_at")
    search_fields = ("usn", "name")
    actions = [reset_login]

    @admin.display(boolean=True, description="Logged in")
    def logged_in(self, obj):
        return bool(obj.active_token)

    def get_urls(self):
        wrap = self.admin_site.admin_view
        custom = [
            path("import/", wrap(self.import_view), name="portal_student_import"),
            path("import/template/", wrap(self.template_view), name="portal_student_template"),
            path("import/dummy/", wrap(self.dummy_view), name="portal_student_dummy"),
        ]
        return custom + super().get_urls()

    def import_view(self, request):
        if request.method == "POST":
            form = ImportForm(request.POST, request.FILES)
            if form.is_valid():
                try:
                    rows = read_rows(form.cleaned_data["file"])
                    created, updated, problems = import_students(rows, form.cleaned_data["exam"])
                except ValueError as e:
                    form.add_error("file", str(e))
                else:
                    messages.success(
                        request,
                        f"Import finished: {created} new, {updated} updated, {len(problems)} rows skipped.",
                    )
                    for p in problems[:10]:
                        messages.warning(request, p)
                    if len(problems) > 10:
                        messages.warning(request, f"...and {len(problems) - 10} more rows skipped.")
                    return redirect("admin:portal_student_changelist")
        else:
            form = ImportForm()

        context = {**self.admin_site.each_context(request), "form": form, "title": "Import students"}
        return render(request, "admin/portal/student/import.html", context)

    def template_view(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(["USN", "Name", "DOB", "Email"])
        ws.append(["2LB01CS001", "Sample Student", datetime.date(2005, 6, 5), "sample@example.com"])
        style_student_sheet(ws)
        return xlsx_response(wb, "student_template.xlsx")

    def dummy_view(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        ws.append(["USN", "Name", "DOB", "Email"])
        for i in range(1, 31):
            ws.append([f"TEST{i:03d}", f"Test Student {i:02d}", datetime.date(2000, 1, 1), ""])
        style_student_sheet(ws)
        return xlsx_response(wb, "dummy_students_30.xlsx")


@admin.register(Exam)
class ExamAdmin(admin.ModelAdmin):
    inlines = [SectionInline]
    filter_horizontal = ("students",)
    actions = [download_results]


@admin.register(Question)
class QuestionAdmin(admin.ModelAdmin):
    form = QuestionAdminForm
    inlines = [OptionInline]
    list_display = ("order", "__str__", "section", "marks")
    list_filter = ("section",)
    fields = ("exam", "section", "order", "text", "marks")

    class Media:
        js = ("portal/admin_question_section.js",)

    def get_urls(self):
        wrap = self.admin_site.admin_view
        custom = [
            path("import/", wrap(self.import_view), name="portal_question_import"),
            path("import/template/", wrap(self.template_view), name="portal_question_template"),
        ]
        return custom + super().get_urls()

    def import_view(self, request):
        if request.method == "POST":
            form = QuestionImportForm(request.POST, request.FILES)
            if form.is_valid():
                try:
                    rows = read_rows(form.cleaned_data["file"])
                    created, updated, problems = import_questions(rows, form.cleaned_data["exam"])
                except ValueError as e:
                    form.add_error("file", str(e))
                else:
                    messages.success(
                        request,
                        f"Import finished: {created} new, {updated} updated, {len(problems)} rows skipped.",
                    )
                    for p in problems[:10]:
                        messages.warning(request, p)
                    if len(problems) > 10:
                        messages.warning(request, f"...and {len(problems) - 10} more rows skipped.")
                    return redirect("admin:portal_question_changelist")
        else:
            form = QuestionImportForm()

        context = {
            **self.admin_site.each_context(request),
            "form": form,
            "title": "Import questions",
        }
        return render(request, "admin/portal/question/import.html", context)

    def template_view(self, request):
        return xlsx_response(question_template_workbook(), "mcq_questions_template.xlsx")

class CodingQuestionStarterCodeInline(admin.TabularInline):
    model = CodingQuestionStarterCode
    extra = 1
    fields = ("language", "code")
    verbose_name = "Starter Code Scaffold"
    verbose_name_plural = "Starter Code Scaffolds"

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name == "language":
            kwargs["widget"] = forms.Select(choices=[("", "-- Select Language --")] + LanguageRegistry.get_all_choices())
        elif db_field.name == "code":
            kwargs["widget"] = forms.Textarea(attrs={"rows": 6, "style": "font-family: Consolas, monospace; width: 100%;"})
        return super().formfield_for_dbfield(db_field, request, **kwargs)


class TestCaseInline(admin.StackedInline):
    model = TestCase
    extra = 0
    fields = (("input_data", "expected_output", "is_sample", "marks"),)
    classes = ("testcase-inline",)
    verbose_name = "Test case"
    verbose_name_plural = "Test cases"

    class Media:
        css = {"all": ("portal/admin.css",)}

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        if db_field.name in ("input_data", "expected_output"):
            kwargs["widget"] = forms.Textarea(attrs={"rows": 5})
        return super().formfield_for_dbfield(db_field, request, **kwargs)


@admin.register(CodingQuestion)
class CodingQuestionAdmin(admin.ModelAdmin):
    form = CodingQuestionAdminForm
    list_display = ("order", "title", "section", "marks")
    list_filter = ("section__exam",)
    inlines = [CodingQuestionStarterCodeInline, TestCaseInline]
    fieldsets = (
        ("Question", {"fields": ("exam", ("section", "order"), "title", "statement")}),
        ("Input and output guidance", {
            "fields": ("input_format", "output_format", "constraints"),
        }),
        ("Sample", {"fields": (("sample_input", "sample_output"),)}),
        ("Scoring", {"fields": (("time_limit_seconds", "marks"),)}),
    )

    class Media:
        js = ("portal/admin_question_section.js",)

    def get_urls(self):
        wrap = self.admin_site.admin_view
        custom = [
            path("import/", wrap(self.import_view), name="portal_codingquestion_import"),
            path("import/template/", wrap(self.template_view), name="portal_codingquestion_template"),
        ]
        return custom + super().get_urls()

    def import_view(self, request):
        if request.method == "POST":
            form = CodingQuestionImportForm(request.POST, request.FILES)
            if form.is_valid():
                try:
                    question_rows, test_case_rows = read_coding_workbook(form.cleaned_data["file"])
                    created, updated, problems = import_coding_questions(
                        question_rows,
                        test_case_rows,
                        form.cleaned_data["exam"],
                    )
                except ValueError as e:
                    form.add_error("file", str(e))
                else:
                    messages.success(
                        request,
                        f"Import finished: {created} new, {updated} updated, {len(problems)} rows skipped.",
                    )
                    for problem in problems[:10]:
                        messages.warning(request, problem)
                    if len(problems) > 10:
                        messages.warning(request, f"...and {len(problems) - 10} more rows skipped.")
                    return redirect("admin:portal_codingquestion_changelist")
        else:
            form = CodingQuestionImportForm()

        context = {
            **self.admin_site.each_context(request),
            "form": form,
            "title": "Import coding questions",
        }
        return render(request, "admin/portal/codingquestion/import.html", context)

    def template_view(self, request):
        return xlsx_response(coding_question_template_workbook(), "coding_questions_template.xlsx")

    def formfield_for_dbfield(self, db_field, request, **kwargs):
        text_rows = {
            "statement": 5,
            "input_format": 5,
            "output_format": 5,
            "constraints": 5,
            "sample_input": 5,
            "sample_output": 5,
        }
        if db_field.name in text_rows:
            kwargs["widget"] = forms.Textarea(attrs={"rows": text_rows[db_field.name]})
        return super().formfield_for_dbfield(db_field, request, **kwargs)


admin.site.site_header = "PyForgeLabs CBT Administration"
admin.site.site_title = "PyForgeLabs CBT Admin"
admin.site.index_title = "CBT Administration Center"


@admin.register(Section)
class SectionAdmin(admin.ModelAdmin):
    list_display = ("title", "exam", "order", "kind", "start_time", "end_time")
    list_filter = ("exam", "kind")
    search_fields = ("title", "exam__title")
    ordering = ("exam", "order")


@admin.register(TestCase)
class TestCaseAdmin(admin.ModelAdmin):
    list_display = ("__str__", "question", "is_sample", "marks")
    list_filter = ("is_sample", "question__section__exam")
    search_fields = ("question__title",)


@admin.register(Attempt)
class AttemptAdmin(admin.ModelAdmin):
    list_display = ("student", "exam", "started_at", "submitted_at")
    list_filter = ("exam", "submitted_at")
    search_fields = ("student__usn", "student__name", "exam__title")
    date_hierarchy = "started_at"


@admin.register(Answer)
class AnswerAdmin(admin.ModelAdmin):
    list_display = ("attempt", "question", "selected_option", "marked_for_review", "updated_at")
    list_filter = ("marked_for_review", "question__section__exam")
    search_fields = ("attempt__student__usn", "question__text")


@admin.register(CodeSubmission)
class CodeSubmissionAdmin(admin.ModelAdmin):
    list_display = ("attempt", "question", "language", "verdict", "passed", "total", "score", "created_at")
    list_filter = ("verdict", "language", "question__section__exam")
    search_fields = ("attempt__student__usn", "question__title")
    date_hierarchy = "created_at"


@admin.register(CodeDraft)
class CodeDraftAdmin(admin.ModelAdmin):
    list_display = ("attempt", "question", "language", "updated_at")
    list_filter = ("language", "question__section__exam")
    search_fields = ("attempt__student__usn", "question__title")


@admin.register(Feedback)
class FeedbackAdmin(admin.ModelAdmin):
    list_display = ("attempt", "experience_rating", "realism_rating", "created_at")
    list_filter = ("experience_rating", "realism_rating")
    search_fields = ("attempt__student__usn", "comments")
    date_hierarchy = "created_at"


@admin.register(ExecutionJob)
class ExecutionJobAdmin(admin.ModelAdmin):
    list_display = ("id", "kind", "status", "attempt", "question", "verdict", "score", "duration_ms", "created_at")
    list_filter = ("status", "kind", "created_at")
    search_fields = ("attempt__student__usn", "question__title", "worker_id")
    date_hierarchy = "created_at"
    readonly_fields = ("id", "created_at", "started_at", "completed_at", "duration_ms")



@admin.register(CodingQuestionStarterCode)
class CodingQuestionStarterCodeAdmin(admin.ModelAdmin):
    list_display = ("coding_question", "language", "updated_at")
    list_filter = ("language", "coding_question__section__exam")
    search_fields = ("coding_question__title", "language", "code")
